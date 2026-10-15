"""DOCX 日报与 XLSX 机舰活动明细表输出，版式对齐官方样例。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION  # noqa: F401
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, Twips
from openpyxl import Workbook
from openpyxl.styles import Font

from . import analysis as A
from .events import Event, parse_day

AIR_COLS = ["日期", "星期", "国家地区", "机型", "批号", "批次", "架次", "起飞\n基地", "呼号", "起飞时间", "降落时间",
            "活动方向", "活动区域", "任务描述", "任务类型", "备注"]
SHIP_COLS = ["日期", "国家地区", "舰船", "批号", "艘次", "活动时间", "活动方向", "活动区域", "任务", "母港", "备注"]
AIR_WIDTH = {"A": 12, "B": 10, "D": 15, "E": 20, "F": 10, "H": 14, "I": 12, "J": 10, "N": 36, "O": 10, "P": 36}
SHIP_WIDTH = {"A": 12, "B": 10, "C": 15, "D": 21, "E": 10, "F": 11, "G": 10, "J": 14, "K": 35}


# ---------------------------------------------------------------- XLSX
def write_xlsx(day: str, events: list[Event], path: Path) -> None:
    ev = A.core(events)
    wb = Workbook()
    wa = wb.active
    wa.title = "飞机"
    ws = wb.create_sheet("舰艇")
    head, body = Font(name="宋体", size=11, bold=True), Font(name="宋体", size=11)
    wa.append(AIR_COLS)
    for e in (x for x in ev if x.kind == "air"):
        wa.append([day, e.weekday, e.country, e.platform, e.batch, 1, e.count, e.base, A.callsign(e), e.start, e.end,
                   e.direction, e.region, A.task_desc(e), e.task, A.remark(e)])
    ws.append(SHIP_COLS)
    for e in (x for x in ev if x.kind == "ship"):
        ws.append([day, e.country, e.platform, e.batch, e.count, f"{e.start}-{e.end}", e.direction, e.region, e.task,
                   e.base, A.remark(e)])
    for sh, widths in ((wa, AIR_WIDTH), (ws, SHIP_WIDTH)):
        for row in sh.iter_rows():
            for c in row:
                c.font = head if c.row == 1 else body
        for col, w in widths.items():
            sh.column_dimensions[col].width = w
        sh.freeze_panes = "A2"
    wb.save(path)


# ---------------------------------------------------------------- DOCX 基础
def _fonts(run, east: str, ascii_: str | None = None, size: float = 16, bold: bool = False) -> None:
    ascii_ = ascii_ or east
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for k, v in (("w:ascii", ascii_), ("w:hAnsi", ascii_), ("w:eastAsia", east), ("w:cs", ascii_)):
        rf.set(qn(k), v)


def _para(doc, text: str, kind: str = "body", align: str = "both", indent: bool = True, line: float = 29):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = Pt(line)
    pf.first_line_indent = Twips(640) if indent else Twips(0)
    if kind in ("title", "h1", "h2"):
        pf.keep_with_next = True
    p.alignment = {"both": 3, "center": 1, "left": 0}[align]
    if text:
        r = p.add_run(text)
        if kind == "title":
            _fonts(r, "方正小标宋简体", size=22, bold=True)
        elif kind == "h1":
            _fonts(r, "黑体", size=16, bold=True)
        elif kind == "h2":
            _fonts(r, "楷体_GB2312", "Times New Roman", size=16, bold=True)
        else:
            _fonts(r, "仿宋_GB2312", "Times New Roman", size=16)
    return p


def _page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _footer_page_number(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = 2
    for text, fld in (("第 ", None), (None, "PAGE"), (" 页", None)):
        r = p.add_run()
        _fonts(r, "仿宋_GB2312", size=9)
        if text:
            r.text = text
            continue
        for t, val in (("begin", None), ("instr", fld), ("end", None)):
            if t == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = val
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), t)
            r._element.append(el)


def _table(doc, header: list[str], rows: list[list[str]], width_twips: int) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.autofit = False
    # Give short identifiers less space and retain room for evidence descriptions.
    widths = ([650, 1250, 1650, 2050, 1300, 1944] if len(header) == 6 else
              [650, 1250, 1700, 650, 650, 1100, 2100, 744] if len(header) == 8 else
              [width_twips] * len(header))
    for column, width in zip(t.columns, widths):
        column.width = Twips(width)
    repeat_header = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat_header)
    for i, row in enumerate([header] + rows):
        # Every record remains readable on one page; long records may still flow
        # when they exceed a page, with the header repeated by Word.
        t.rows[i]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Twips(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = 1 if i == 0 or j == 0 else 0
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            parts = str(val).split("\n")
            r = p.add_run(parts[0])
            for extra in parts[1:]:
                r.add_break()
                r.add_text(extra)
            _fonts(r, "黑体" if i == 0 else "仿宋_GB2312", "黑体" if i == 0 else "Times New Roman", size=10.5)


# ---------------------------------------------------------------- DOCX 日报
def write_docx(day: str, events: list[Event], prev: list[Event] | None,
               history: list[tuple[str, list[Event]]], charts: list[Path], path: Path) -> None:
    d = parse_day(day)
    date = f"{d.year}年{d.month}月{d.day}日"
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(11906), Twips(16838)
    sec.top_margin, sec.bottom_margin = Twips(2098), Twips(1985)
    sec.left_margin, sec.right_margin = Twips(1588), Twips(1474)
    _footer_page_number(sec)

    _para(doc, f"{d.month}月{d.day}日飞机舰船活动情况", "title", "center", indent=False, line=22)
    _para(doc, A.overview(day, events, prev))

    _para(doc, "一、重要军事动向", "h1")
    for s in A.section_important(events):
        _para(doc, s)

    _para(doc, "二、重点态势分析", "h1")
    focus, focus_title = A.section_focus(history, day, events)
    for s in focus:
        _para(doc, s)

    _para(doc, "三、外（台）机舰活动情况", "h1")
    rows = A.section_countries(events)
    if not rows:
        _para(doc, "当日未掌握外（台）机舰在东海、南海、黄海方向的可信活动。")
    for head, text in rows:
        _para(doc, head, "h2")
        _para(doc, text)

    _para(doc, "四、美重要武器平台部署情况", "h1")
    _para(doc, A.section_us(events))

    _para(doc, "五、次日关注重点", "h1")
    _para(doc, A.section_next(events, focus_title))

    titles = ["舰机侦巡数量曲线图", f"{date}中国主要机舰活动示意图", f"{date}美国主要机舰活动示意图",
              f"{date}日、加等其他国家主要机舰活动示意图", f"{date}主要海军基地舰艇联合点名报告表",
              f"{date}主要空军基地飞机联合点名报告表"]
    for i, t in enumerate(titles, 1):
        _para(doc, ("附件：" if i == 1 else "") + f"{i}．{t}", indent=True)

    content_w = 11906 - 1588 - 1474
    for i, t in enumerate(titles, 1):
        _page_break(doc)
        _para(doc, f"附件{i}", "h1", "left", indent=False)
        _para(doc, t, "h1", "center", indent=False)
        if i <= 4:
            p = _para(doc, "", align="center", indent=False, line=12)
            p.paragraph_format.line_spacing = 1.0
            p.add_run().add_picture(str(charts[i - 1]), width=Emu(content_w * 635))
        elif i == 5:
            _table(doc, ["国家/\n地区", "基地", "常态部署舰艇", "最新侦察情况", "侦察手段", "备注"],
                   A.naval_table(events), 1474)
        else:
            _table(doc, ["国家/\n地区", "机场", "机型", "常驻数量", "现驻\n数量", "驻停\n位置", "执行作战\n勤务情况", "备注"],
                   A.air_table(events), 1105)
    doc.save(path)
