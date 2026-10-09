from __future__ import annotations

import json
import re
from collections import defaultdict
from dataclasses import dataclass, asdict, field
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook


@dataclass
class EventRecord:
    date: str
    source: str
    title: str
    platform: str
    task: str
    region: str
    action: str
    raw_text: str
    confidence: float = 0.0
    source_set: List[str] = field(default_factory=list)


class CompetitionPipeline:
    """A working pipeline for multi-source military event extraction and e2e reporting."""

    def __init__(self, data_root: Path):
        self.data_root = Path(data_root)
        self.platform_keywords = [
            "歼-16D", "歼-35", "苏-35S", "苏-30SM", "F-15SG", "AW-159", "P-3CK", "E-7A",
            "E-2C", "RC-135V", "轰-6K", "空警-500", "903A", "054B", "054A", "039A", "056A",
            "福建号", "山东号", "071", "秋月级", "大邱级", "村雨号", "新奥尔良号", "阿利·伯克级",
            "瓦良格号", "独岛级", "米利厄斯号", "莱库号", "F-35B", "台风", "运-8", "苏-30MK2"
        ]
        self.task_keywords = [
            "反潜巡逻", "战备警巡", "电子对抗训练", "对海突击训练", "联合巡航",
            "空中警戒", "海上护航", "远海训练", "巡逻", "起降训练", "两栖作战训练",
            "联合搜救训练", "侦察", "海上侦察", "编队警戒部署", "实弹射击", "持续活动"
        ]
        self.region_keywords = [
            "东海", "南海", "黄海", "阿拉伯海", "塔斯曼海", "孟加拉湾", "巴伦支海",
            "安达曼海", "渚碧礁", "仁爱礁", "礼乐滩", "琼州海峡", "温州外海", "上海外海",
            "钓鱼岛", "中太平洋", "菲律宾海", "鄂霍次克海", "印度洋", "西沙群岛", "南沙群岛",
            "中部", "北部", "南部", "东海南部", "中太平洋海域", "中部海域"
        ]
        self.platform_aliases = {
            "054b": "054b",
            "054b型导弹护卫舰": "054b",
            "054b型": "054b",
            "054a": "054a",
            "p8a": "p8a",
            "p-8a": "p8a",
            "p-8a海上侦察机": "p8a",
            "rc135v": "rc135v",
            "f35b": "f35b",
            "f-35b": "f35b",
        }
        self.task_aliases = {
            "反潜": "反潜巡逻",
            "反潜巡逻": "反潜巡逻",
            "巡航": "联合巡航",
            "联合巡航": "联合巡航",
            "侦察": "侦察",
            "海上侦察": "侦察",
            "警戒": "空中警戒",
            "空中警戒": "空中警戒",
            "活动": "持续活动",
            "持续活动": "持续活动",
        }
        self.region_aliases = {
            "东海海域": "东海",
            "东海北部": "东海",
            "东海中部": "东海",
            "东海南部": "东海",
            "钓鱼岛海域": "东海",
            "赤尾屿海域": "东海",
            "上海外海": "东海",
            "温州外海": "东海",
            "长江口外海": "东海",
            "冲绳海槽": "东海",
            "南海海域": "南海",
            "南海北部": "南海",
            "南海中部": "南海",
            "南海南部": "南海",
            "北部湾": "南海",
            "琼州海峡": "南海",
            "西沙群岛海域": "南海",
            "中沙群岛海域": "南海",
            "南沙群岛海域": "南海",
            "黄岩岛海域": "南海",
            "仁爱礁附近海域": "南海",
            "礼乐滩海域": "南海",
            "渚碧礁附近海域": "南海",
            "永暑礁附近海域": "南海",
            "美济礁附近海域": "南海",
            "曾母暗沙海域": "南海",
            "黄海海域": "黄海",
            "北黄海": "黄海",
            "南黄海": "黄海",
            "黄海中部": "黄海",
            "青岛外海": "黄海",
            "成山头外海": "黄海",
            "海州湾外海": "黄海",
            "连云港外海": "黄海",
            "济州岛西部海域": "黄海",
            "渚碧礁": "南海",
            "仁爱礁": "南海",
            "礼乐滩": "南海",
            "西沙群岛": "南海",
            "南沙群岛": "南海",
            "中部海域": "南海",
            "菲律宾海": "其他",
            "中太平洋": "其他",
            "巴士海峡": "其他",
            "台湾海峡": "其他",
            "地中海": "其他",
            "黑海": "其他",
            "波斯湾": "其他",
            "孟加拉湾": "其他",
            "阿拉伯海": "其他",
            "塔斯曼海": "其他",
            "印度洋": "其他",
        }

    def run_for_date(self, date_str: str) -> Dict[str, Any]:
        date_dir = self.data_root / date_str
        if not date_dir.exists():
            raise FileNotFoundError(f"Date folder not found: {date_dir}")

        text_records = []
        for source_name in ["开源军情", "网上舆论"]:
            source_dir = date_dir / source_name
            if source_dir.exists():
                text_records.extend(self._parse_txt_source(source_dir, source_name, date_str))

        radar_records = self._parse_radar_records(date_dir, date_str)
        merged_events = self._merge_text_and_radar(text_records, radar_records)
        credible_events = self._filter_credible_events(merged_events)

        summary = {
            "date": date_str,
            "text_record_count": len(text_records),
            "radar_record_count": len(radar_records),
            "merged_event_count": len(merged_events),
            "credible_event_count": len(credible_events),
            "events": [asdict(evt) for evt in credible_events[:20]],
        }
        return summary

    def _parse_txt_source(self, source_dir: Path, source_name: str, date_str: str) -> List[EventRecord]:
        records: List[EventRecord] = []
        for txt_path in sorted(source_dir.glob("*.txt")):
            text = txt_path.read_text(encoding="utf-8-sig")
            title = self._extract_title(text)
            platform = self._extract_platform(title, text)
            task = self._extract_task(title, text)
            region = self._extract_region(title, text)
            records.append(
                EventRecord(
                    date=date_str,
                    source=source_name,
                    title=title,
                    platform=platform,
                    task=task,
                    region=region,
                    action=self._extract_action(text),
                    raw_text=text,
                    confidence=0.7,
                    source_set=[source_name],
                )
            )
        return records

    def _parse_radar_records(self, date_dir: Path, date_str: str) -> List[EventRecord]:
        records: List[EventRecord] = []
        for xlsx_path in sorted(date_dir.glob("*.xlsx")):
            try:
                workbook = load_workbook(xlsx_path, read_only=True)
            except Exception:
                continue
            ws = workbook[workbook.sheetnames[0]]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(h).strip() for h in rows[0]]
            for row in rows[1:]:
                if not any(v is not None and str(v).strip() for v in row):
                    continue
                row_dict = {headers[idx]: row[idx] if idx < len(row) else None for idx in range(len(headers))}
                platform = str(row_dict.get("目标类型") or row_dict.get("目标名称") or "未知平台").strip()
                task = str(row_dict.get("任务") or "未知任务").strip()
                region = self._extract_region(
                    str(row_dict.get("到达地") or "") + " " +
                    str(row_dict.get("起飞机场") or "") + " " +
                    str(row_dict.get("国家地区") or ""),
                    ""
                )
                records.append(
                    EventRecord(
                        date=date_str,
                        source="雷达",
                        title=f"雷达记录: {platform}",
                        platform=platform,
                        task=task,
                        region=region,
                        action=self._extract_action(str(task)),
                        raw_text=json.dumps(row_dict, ensure_ascii=False),
                        confidence=0.85,
                        source_set=["雷达"],
                    )
                )
        return records

    def _event_group_key(self, event: EventRecord) -> tuple:
        base_key = (
            self._norm(event.date),
            self._canonical_platform(event.platform),
            self._normalize_task(event.task),
        )
        region = self._normalize_region(event.region)
        if region and region != '未知区域':
            return base_key + (region,)
        return base_key

    def _event_identity_key(self, event: EventRecord) -> tuple:
        region = self._normalize_region(event.region)
        return (
            self._norm(event.date),
            self._canonical_platform(event.platform),
            self._normalize_task(event.task),
            region if region and region != '未知区域' else '其他',
        )

    def _source_strength(self, record: EventRecord) -> float:
        strength = float(record.confidence)
        if record.source == '雷达':
            strength += 0.1
        if record.source in {'开源军情', '网上舆论'}:
            strength += 0.05
        return strength

    def _merge_text_and_radar(self, text_records: List[EventRecord], radar_records: List[EventRecord]) -> List[EventRecord]:
        """Merge same-event candidates only when their key semantic fields line up."""
        merged: List[EventRecord] = []
        grouped: Dict[tuple, List[EventRecord]] = defaultdict(list)

        all_records = text_records + radar_records
        for record in all_records:
            grouped[self._event_identity_key(record)].append(record)

        for records in grouped.values():
            if not records:
                continue

            actions = {self._normalize_action(r.action) for r in records if r.action}
            if len(actions) > 1 and {'进入', '离开'} <= actions:
                continue

            core_regions = [r.region for r in records if self._normalize_region(r.region) in {'东海', '南海', '黄海'}]
            if core_regions:
                preferred_region = max(core_regions, key=lambda region: len(str(region)))
            else:
                preferred_region = max(
                    [r.region for r in records if r.region and r.region != '未知区域'],
                    key=lambda region: (1 if '海' in str(region) else 0, len(str(region))),
                    default='未知区域',
                )

            best = max(
                records,
                key=lambda r: (
                    1 if self._normalize_region(r.region) in {'东海', '南海', '黄海'} else 0,
                    len(r.source_set) if r.source_set else 1,
                    self._source_strength(r),
                    r.confidence,
                    1 if r.source == '雷达' else 0,
                    len(str(r.title)),
                ),
            )
            best.region = preferred_region if preferred_region else best.region

            text_candidates = [r for r in records if r.platform and r.platform != '未知平台']
            if text_candidates:
                preferred = max(
                    text_candidates,
                    key=lambda r: (
                        1 if r.source in {'开源军情', '网上舆论'} else 0,
                        len(str(r.platform)),
                        r.confidence,
                    ),
                )
                best.platform = preferred.platform

            if best.task in {'未知任务', ''}:
                non_unknown = [r.task for r in records if r.task and r.task != '未知任务']
                if non_unknown:
                    best.task = non_unknown[0]

            if not best.title or best.title == '未知标题':
                best.title = f"{best.platform}在{best.region}开展{best.task}活动"

            best.source_set = sorted({s for r in records for s in (r.source_set if r.source_set else [r.source])})
            merged.append(best)

        return sorted(merged, key=lambda x: (x.date, x.region, x.task, x.platform))

    def _filter_credible_events(self, events: List[EventRecord]) -> List[EventRecord]:
        """Official contest rule: keep only cross-source-supported facts in the core sea regions."""
        grouped: Dict[tuple, List[EventRecord]] = defaultdict(list)
        for event in events:
            grouped[self._event_identity_key(event)].append(event)

        credible: List[EventRecord] = []
        for records in grouped.values():
            sources = set()
            for r in records:
                if r.source_set:
                    sources.update(r.source_set)
                else:
                    sources.add(r.source)

            if len(sources) < 2:
                continue

            region_values = [self._normalize_region(r.region) for r in records if r.region]
            if not any(region in {'东海', '南海', '黄海'} for region in region_values):
                continue

            actions = {self._normalize_action(r.action) for r in records if r.action}
            if len(actions) > 1 and {'进入', '离开'} <= actions:
                continue

            if sources == {'雷达'}:
                continue

            chosen = max(
                records,
                key=lambda r: (
                    1 if self._normalize_region(r.region) in {'东海', '南海', '黄海'} else 0,
                    len(r.source_set) if r.source_set else 1,
                    self._source_strength(r),
                    r.confidence,
                    1 if r.source == '雷达' else 0,
                    len(str(r.title)),
                ),
            )
            chosen.source_set = sorted(sources)
            chosen.confidence = 0.9 if len(sources) >= 3 else 0.8
            credible.append(chosen)

        return sorted(credible, key=lambda x: (-x.confidence, x.date, x.region, x.task, x.platform))

    def _canonical_platform(self, platform: str) -> str:
        s = self._norm(platform)
        if not s:
            return ""

        normalized = self.platform_aliases.get(s)
        if normalized:
            return normalized

        for token in sorted(self.platform_keywords, key=lambda x: (-len(x), x)):
            key = self._norm(token)
            if key and key in s:
                return self.platform_aliases.get(key, key)

        suffixes = [
            '多用途战斗机', '战斗机', '轰炸机', '预警机', '巡逻机', '反潜直升机', '直升机',
            '导弹护卫舰', '护卫舰', '综合补给舰', '驱逐舰', '短距起降战斗机', '反潜巡逻机',
            '舰载机', '舰艇', '飞机', '舰', '机'
        ]
        candidate = s
        for suffix in suffixes:
            suffix_key = self._norm(suffix)
            if candidate.endswith(suffix_key):
                candidate = candidate[:-len(suffix_key)]
        if candidate and candidate != s:
            return candidate
        return s

    def analyze_cross_day_trend(self, date_str: str, days_back: int = 7) -> Dict[str, Any]:
        """Collect the recent activity history for the target date and summarize focus-event trend."""
        date_dir = self.data_root / date_str
        if not date_dir.exists():
            raise FileNotFoundError(f"Date folder not found: {date_dir}")

        all_dates = sorted(
            [p.name for p in self.data_root.iterdir() if p.is_dir() and re.fullmatch(r"\d{4}\.\d{1,2}\.\d{1,2}", p.name)],
            key=lambda x: tuple(int(part) for part in x.split('.')),
        )
        target_index = all_dates.index(date_str) if date_str in all_dates else len(all_dates) - 1
        recent_dates = all_dates[max(0, target_index - days_back + 1):target_index + 1]

        history: List[EventRecord] = []
        for d in recent_dates:
            try:
                summary = self.run_for_date(d)
            except FileNotFoundError:
                continue
            for ev in summary.get('events', []):
                record = EventRecord(**ev)
                if any(region in record.region for region in ['东海', '南海', '黄海']):
                    history.append(record)

        by_day = defaultdict(int)
        by_region = defaultdict(int)
        by_platform = defaultdict(int)
        for ev in history:
            by_day[ev.date] += 1
            by_region[ev.region] += 1
            by_platform[ev.platform] += 1

        focus_event = max(history, key=lambda r: (r.confidence, 1 if r.source == '雷达' else 0), default=None)
        return {
            "date": date_str,
            "history_dates": recent_dates,
            "total_events": len(history),
            "by_day": dict(sorted(by_day.items())),
            "by_region": dict(sorted(by_region.items())),
            "by_platform": dict(sorted(by_platform.items(), key=lambda kv: (-kv[1], kv[0]))[:10]),
            "focus_event": asdict(focus_event) if focus_event else None,
            "trend_summary": (
                "持续活动" if len(by_day) > 2 and max(by_day.values()) >= 2 else "局部波动"
            ),
        }

    def _extract_title(self, text: str) -> str:
        match = re.search(r"标题：\s*(.+)", text, flags=re.M)
        if match:
            return match.group(1).strip()
        return "未知标题"

    def _extract_action(self, text: str) -> str:
        if "进入" in text:
            return "进入"
        if "离开" in text:
            return "离开"
        if "活动" in text or "执行" in text or "组织" in text:
            return "活动"
        return "未知"

    def _extract_platform(self, title: str, text: str) -> str:
        combined = title + "\n" + text
        for token in self.platform_keywords:
            if token in combined:
                return token
        m = re.search(r"([A-Za-z0-9-]+(?:[型级号]|战斗机|驱逐舰|护卫舰|舰载机|轰炸机|预警机|直升机|潜艇|巡逻机|航母|登陆舰|舰艇))", combined)
        if m:
            return m.group(1)
        return "未知平台"

    def _extract_task(self, title: str, text: str) -> str:
        combined = title + "\n" + text
        for token in self.task_keywords:
            if token in combined:
                return token
        return "未知任务"

    def _extract_region(self, title: str, text: str) -> str:
        combined = (title + "\n" + text).strip()
        for token in self.region_keywords:
            if token in combined:
                return token
        return "未知区域"

    def _normalize_task(self, task: str) -> str:
        s = self._norm(task)
        if not s:
            return "未知任务"
        for key, normalized in self.task_aliases.items():
            if s == self._norm(key):
                return normalized
        if "反潜" in s:
            return "反潜巡逻"
        if "巡逻" in s:
            return "反潜巡逻"
        if "侦察" in s:
            return "侦察"
        if "警戒" in s:
            return "空中警戒"
        if "活动" in s:
            return "持续活动"
        return s

    def _normalize_region(self, region: str) -> str:
        s = self._norm(region)
        if not s:
            return "未知区域"
        for key, normalized in self.region_aliases.items():
            if s == self._norm(key):
                return normalized
        if "东海" in s:
            return "东海"
        if "南海" in s:
            return "南海"
        if "黄海" in s:
            return "黄海"
        if "渤海" in s or "莱州湾" in s or "海州湾" in s:
            return "黄海"
        if "菲律宾海" in s or "中太平洋" in s or "台湾海峡" in s or "巴士海峡" in s:
            return "其他"
        return s

    def _normalize_action(self, action: str) -> str:
        s = self._norm(action)
        if "进入" in s:
            return "进入"
        if "离开" in s:
            return "离开"
        if "活动" in s or "执行" in s or "组织" in s:
            return "活动"
        return "未知"

    def _norm(self, value: Optional[str]) -> str:
        if value is None:
            return ""
        return re.sub(r"\s+", "", str(value)).lower()


def main() -> None:
    base_dir = Path(__file__).resolve().parents[2] / "数据"
    pipeline = CompetitionPipeline(base_dir)
    summary = pipeline.run_for_date("2025.5.1")
    out_dir = Path(__file__).resolve().parents[2] / "outputs"
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / "2025.5.1_summary.json"
    out_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "date": summary["date"],
        "text_record_count": summary["text_record_count"],
        "radar_record_count": summary["radar_record_count"],
        "merged_event_count": summary["merged_event_count"],
        "credible_event_count": summary["credible_event_count"],
        "sample_event": summary["events"][0] if summary["events"] else None,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
