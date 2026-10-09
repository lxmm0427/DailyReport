"""以官方样例(2025.10.20)为基准的回归测试，以及提交物格式校验。"""
import sys
from pathlib import Path

import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

from dailyreport import analysis as A  # noqa: E402
from dailyreport.pipeline import history_until, process_day  # noqa: E402
from dailyreport.textparse import Mention, extract  # noqa: E402

SAMPLE = ROOT / "提交示例文件"
DAY = "2025.10.20"
GT_XLSX = SAMPLE / "日报" / DAY / f"{DAY}机舰活动明细.xlsx"


@pytest.fixture(scope="module")
def sample_day():
    return process_day(SAMPLE / "数据", DAY)


def _gt():
    return pd.concat([pd.read_excel(GT_XLSX, sheet_name=s) for s in ("飞机", "舰艇")])


def test_trusted_batches_match_sample(sample_day):
    gt = _gt()
    pred = {e.batch: e for e in sample_day.events if e.direction != "其他"}
    assert set(pred) == set(gt["批号"])
    for b, r in zip(gt["批号"], gt["活动区域"]):
        assert pred[b].region == r


def test_overview_matches_sample(sample_day):
    prev = process_day(SAMPLE / "数据", "2025.10.19")
    text = A.overview(DAY, sample_day.events, prev.events)
    assert text.startswith("10月20日，共掌握外（台）各型飞机10批11架次、舰船15艘位当面活动，"
                           "与1日前相比，飞机减少6架次、舰船增加1艘")


def test_hourly_counts_match_sample_chart(sample_day):
    air, ship = A.hourly_counts(sample_day.events)
    assert air == [0, 0, 1, 2, 1, 1, 2, 1, 1, 2, 0, 1, 2, 1, 3, 4, 3, 1, 2, 1, 1, 2, 2, 0]
    assert ship == [0, 0, 1, 2, 1, 1, 5, 2, 5, 6, 5, 1, 2, 1, 1, 1, 1, 3, 3, 2, 1, 2, 1, 0]


def test_history_never_reads_after_node():
    hist = history_until(SAMPLE / "数据", "2025.10.5")
    assert [h.day for h in hist] == ["2025.10.1", "2025.10.2", "2025.10.3", "2025.10.4", "2025.10.5"]


@pytest.mark.parametrize("body,field,value", [
    ("2025年6月16日，英国RC-135W电子侦察机在东海方向执行电子侦察任务，活动时段为11时21分至13时46分。", "start", "1121"),
    ("据俄罗斯卫星通讯社报道，2025年5月20日，中国一架歼-11B战斗机在印度洋活动。", "country", "中国"),
    ("2025年5月25日，日本P-3C反潜巡逻机在宁波外海活动，执行海上侦察巡逻。", "task", "海上侦察巡逻"),
    ("2025年5月25日，加拿大一架“CP-140‘极光’海上巡逻机”在黄海中部海域执行情报监视侦察任务。", "task", "情报监视侦察"),
    ("2025年6月1日，日本海上自卫队P-1反潜巡逻机在东海南部活动。", "region", "东海南部"),
    ("2025.10.20，新西兰P-8A“海神”导弹巡洋舰在青岛外海开展行动，涉及1艘。", "kind_conflict", True),
])
def test_text_extraction(body, field, value):
    m = extract(Mention(source="军情", day="2025.6.1", file="x", title="", body=body))
    assert getattr(m, field) == value


@pytest.fixture(scope="module")
def generated(tmp_path_factory):
    import main
    out = tmp_path_factory.mktemp("out") / "日报"
    main.run_node(SAMPLE / "数据", out, DAY, {})
    return out / DAY


def test_submission_contract(generated):
    files = {p.relative_to(generated).as_posix() for p in generated.rglob("*") if p.is_file()}
    assert files == {f"{DAY}日报.docx", f"{DAY}机舰活动明细.xlsx",
                     "附件图表/01_舰机侦巡数量曲线图.png", "附件图表/02_中国主要机舰活动示意图.png",
                     "附件图表/03_美国主要机舰活动示意图.png", "附件图表/04_其他国家主要机舰活动示意图.png"}
    gt, ours = load_workbook(GT_XLSX), load_workbook(generated / f"{DAY}机舰活动明细.xlsx")
    assert ours.sheetnames == gt.sheetnames == ["飞机", "舰艇"]
    for s in ("飞机", "舰艇"):
        assert [c.value for c in ours[s][1]] == [c.value for c in gt[s][1]]
        assert [type(c.value) for c in ours[s][2]] == [type(c.value) for c in gt[s][2]]
    text = [p.text for p in Document(generated / f"{DAY}日报.docx").paragraphs]
    assert text[0] == "10月20日飞机舰船活动情况"
    for h in ("一、重要军事动向", "二、重点态势分析", "三、外（台）机舰活动情况", "四、美重要武器平台部署情况", "五、次日关注重点"):
        assert h in text
    for p in (generated / "附件图表").glob("*.png"):
        Image.open(p).verify()
